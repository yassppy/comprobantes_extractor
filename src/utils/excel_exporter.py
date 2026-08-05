"""Exportador de comprobantes a Excel (REQ-10).

R1: Exportar comprobantes del período (mes/año) seleccionado.
R2: Generar archivo .xlsx.
R3: Incluir empresa, fecha, serie, número, proveedor/cliente, subtotal,
    IGV, total, código de categoría y nombre de categoría.
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

# ── Paleta de colores ──────────────────────────────────────────────────────────
_COLOR_HEADER_BG   = "1F4E79"   # azul oscuro
_COLOR_HEADER_FG   = "FFFFFF"   # blanco
_COLOR_SUBHEAD_BG  = "2E75B6"   # azul medio
_COLOR_ALT_ROW     = "EBF3FB"   # azul muy claro (filas pares)
_COLOR_PURCHASE    = "FFF2CC"   # amarillo suave (compras)
_COLOR_SALE        = "E2EFDA"   # verde suave (ventas)
_COLOR_TOTAL_BG    = "D6E4F0"   # azul claro (fila de totales)
_COLOR_BORDER      = "B8CCE4"

_THIN = Side(style="thin", color=_COLOR_BORDER)
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# Columnas del Excel según REQ-10 R3
_COLUMNS: list[tuple[str, str, int]] = [
    # (campo_en_dict,       encabezado_excel,           ancho)
    ("empresa",             "Empresa",                   30),
    ("empresa_ruc",         "RUC Empresa",               14),
    ("document_type",       "Tipo Operación",            14),
    ("invoice_type",        "Tipo Comprobante",          24),
    ("series",              "Serie",                      8),
    ("number",              "Número",                    12),
    ("issue_date",          "Fecha Emisión",             14),
    ("currency",            "Moneda",                     9),
    ("partner_doc",         "RUC / DNI Socio",           16),
    ("partner_name",        "Proveedor / Cliente",       35),
    ("description",         "Descripción",               40),
    ("subtotal",            "Subtotal",                  12),
    ("igv",                 "IGV",                       10),
    ("total",               "Total",                     12),
    ("category_code",       "Cód. Categoría",            16),
    ("category_name",       "Nombre Categoría",          26),
    ("status",              "Estado",                    10),
]


def _header_font(bold: bool = True, size: int = 10, color: str = _COLOR_HEADER_FG) -> Font:
    return Font(name="Calibri", bold=bold, size=size, color=color)


def _cell_font(bold: bool = False, size: int = 10) -> Font:
    return Font(name="Calibri", bold=bold, size=size)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _fmt_date(value: Any) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value) if value else ""


def _fmt_money(value: Any) -> float | str:
    if value is None:
        return ""
    try:
        return float(value)
    except (ValueError, TypeError):
        return str(value)


def _doc_type_label(value: str | None) -> str:
    return {"PURCHASE": "Compra", "SALE": "Venta"}.get(value or "", value or "")


def build_excel(
    rows: list[dict[str, Any]],
    company_name: str,
    month: int,
    year: int,
    document_type: str = "ALL",
) -> bytes:
    """
    Construye el archivo Excel en memoria y retorna los bytes.

    Args:
        rows:          Registros de la BD (resultado de export_documents).
        company_name:  Nombre de la empresa (para el título).
        month:         Mes del período exportado.
        year:          Año del período exportado.
        document_type: 'PURCHASE' | 'SALE' | 'ALL'.

    Returns:
        Bytes del archivo .xlsx listo para descargar.
    """
    _MONTH_NAMES = [
        "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
        "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
    ]

    wb = Workbook()
    ws = wb.active
    ws.title = f"{_MONTH_NAMES[month - 1]} {year}"

    n_cols = len(_COLUMNS)

    # ── Fila 1: Título principal ───────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    title_cell = ws.cell(row=1, column=1)
    tipo_label = {"PURCHASE": "Compras", "SALE": "Ventas", "ALL": "Compras y Ventas"}.get(
        document_type, "Comprobantes"
    )
    title_cell.value = f"REGISTRO DE {tipo_label.upper()} — {company_name.upper()}"
    title_cell.font = Font(name="Calibri", bold=True, size=13, color=_COLOR_HEADER_FG)
    title_cell.fill = _fill(_COLOR_HEADER_BG)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # ── Fila 2: Subtítulo período ──────────────────────────────────────────
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    sub_cell = ws.cell(row=2, column=1)
    sub_cell.value = f"Período: {_MONTH_NAMES[month - 1]} {year}   |   Total registros: {len(rows)}"
    sub_cell.font = Font(name="Calibri", bold=False, size=10, color=_COLOR_HEADER_FG)
    sub_cell.fill = _fill(_COLOR_SUBHEAD_BG)
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 16

    # ── Fila 3: vacía de separación ────────────────────────────────────────
    ws.row_dimensions[3].height = 6

    # ── Fila 4: Encabezados de columnas ───────────────────────────────────
    HEADER_ROW = 4
    for col_idx, (_, header, width) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col_idx, value=header)
        cell.font = _header_font()
        cell.fill = _fill(_COLOR_HEADER_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[HEADER_ROW].height = 28

    # ── Filas de datos ─────────────────────────────────────────────────────
    subtotal_sum = 0.0
    igv_sum = 0.0
    total_sum = 0.0

    for row_idx, doc in enumerate(rows, start=HEADER_ROW + 1):
        doc_type = doc.get("document_type") or ""
        is_even = (row_idx % 2 == 0)

        if doc_type == "PURCHASE":
            row_fill = _fill(_COLOR_PURCHASE)
        elif doc_type == "SALE":
            row_fill = _fill(_COLOR_SALE)
        else:
            row_fill = _fill(_COLOR_ALT_ROW) if is_even else None

        for col_idx, (field, _, _w) in enumerate(_COLUMNS, start=1):
            raw_val = doc.get(field)

            # Formatear según tipo de campo
            if field == "issue_date":
                value = _fmt_date(raw_val)
            elif field in ("subtotal", "igv", "total"):
                value = _fmt_money(raw_val)
            elif field == "document_type":
                value = _doc_type_label(raw_val)
            elif field == "status":
                value = {"PROCESSED": "Procesado", "ERROR": "Error", "PENDING": "Pendiente"}.get(
                    raw_val or "", raw_val or ""
                )
            else:
                value = raw_val if raw_val is not None else ""

            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = _cell_font()
            cell.border = _BORDER

            # Alineación
            if field in ("subtotal", "igv", "total"):
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = '#,##0.00'
            elif field in ("series", "number", "currency", "status"):
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left", wrap_text=False)

            if row_fill:
                cell.fill = row_fill

        # Acumular totales
        sub = doc.get("subtotal")
        igv = doc.get("igv")
        tot = doc.get("total")
        if sub is not None:
            try:
                subtotal_sum += float(sub)
            except (ValueError, TypeError):
                pass
        if igv is not None:
            try:
                igv_sum += float(igv)
            except (ValueError, TypeError):
                pass
        if tot is not None:
            try:
                total_sum += float(tot)
            except (ValueError, TypeError):
                pass

    # ── Fila de totales ────────────────────────────────────────────────────
    total_row = HEADER_ROW + len(rows) + 1
    ws.merge_cells(
        start_row=total_row, start_column=1,
        end_row=total_row, end_column=11,  # hasta "Descripción"
    )
    label_cell = ws.cell(row=total_row, column=1, value="TOTALES")
    label_cell.font = Font(name="Calibri", bold=True, size=10, color="000000")
    label_cell.fill = _fill(_COLOR_TOTAL_BG)
    label_cell.alignment = Alignment(horizontal="right", vertical="center")
    label_cell.border = _BORDER

    # Columnas de montos: subtotal=col 12, igv=13, total=14
    money_totals = {12: subtotal_sum, 13: igv_sum, 14: total_sum}
    for col_idx in range(1, n_cols + 1):
        cell = ws.cell(row=total_row, column=col_idx)
        if col_idx in money_totals:
            cell.value = money_totals[col_idx]
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal="right")
        cell.font = Font(name="Calibri", bold=True, size=10)
        cell.fill = _fill(_COLOR_TOTAL_BG)
        cell.border = _BORDER
    ws.row_dimensions[total_row].height = 18

    # ── Freeze panes (encabezado fijo al hacer scroll) ─────────────────────
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1)

    # ── Auto-filter ───────────────────────────────────────────────────────
    ws.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(n_cols)}{HEADER_ROW + len(rows)}"
    )

    # ── Guardar en buffer de memoria ──────────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
